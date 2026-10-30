"""Export a result set (columns + rows) or a conversation as CSV, XLSX, JSON, TXT or Markdown."""

from __future__ import annotations

import csv
import io
import json
from typing import Any

FORMATS = ("csv", "xlsx", "json", "txt", "md")
MEDIA = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
    "txt": "text/plain; charset=utf-8",
    "md": "text/markdown; charset=utf-8",
}


def export_rows(
    columns: list[str],
    rows: list[list[Any]],
    fmt: str,
    *,
    title: str = "",
    sql: str = "",
    meta: dict[str, Any] | None = None,
) -> bytes:
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\n")
        w.writerow(columns)
        w.writerows(rows)
        return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly
    if fmt == "json":
        payload = {
            "title": title,
            "columns": columns,
            "rows": [dict(zip(columns, r)) for r in rows],
            "sql": sql,
            **(meta or {}),
        }
        return json.dumps(payload, indent=2, default=str).encode("utf-8")
    if fmt == "txt":
        return _fixed_width(columns, rows, title, sql).encode("utf-8")
    if fmt == "md":
        return _markdown(columns, rows, title, sql).encode("utf-8")
    if fmt == "xlsx":
        return _xlsx(columns, rows, title, sql, meta or {})
    raise ValueError(f"unsupported format {fmt}; use one of {', '.join(FORMATS)}")


def _fixed_width(columns: list[str], rows: list[list[Any]], title: str, sql: str) -> str:
    widths = [
        max(len(str(c)), *(len(_cell(r[i])) for r in rows)) if rows else len(str(c)) for i, c in enumerate(columns)
    ]
    lines = []
    if title:
        lines += [title, "=" * len(title), ""]
    lines.append("  ".join(str(c).ljust(widths[i]) for i, c in enumerate(columns)))
    lines.append("  ".join("-" * w for w in widths))
    for r in rows:
        lines.append("  ".join(_cell(v).ljust(widths[i]) for i, v in enumerate(r)))
    lines.append(f"\n{len(rows)} rows")
    if sql:
        lines += ["", "SQL:", sql]
    return "\n".join(lines) + "\n"


def _markdown(columns: list[str], rows: list[list[Any]], title: str, sql: str) -> str:
    lines = [f"# {title}", ""] if title else []
    lines.append("| " + " | ".join(columns) + " |")
    lines.append("|" + "---|" * len(columns))
    for r in rows:
        lines.append("| " + " | ".join(_cell(v).replace("|", "\\|") for v in r) + " |")
    lines.append(f"\n{len(rows)} rows")
    if sql:
        lines += ["", "```sql", sql, "```"]
    return "\n".join(lines) + "\n"


def _xlsx(columns: list[str], rows: list[list[Any]], title: str, sql: str, meta: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Result"
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F6FED")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([_xlsx_value(v) for v in r])
    for i, col in enumerate(columns, 1):
        width = max(len(str(col)), *(len(_cell(r[i - 1])) for r in rows[:500])) if rows else len(str(col))
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
        if rows and all(isinstance(r[i - 1], (int, float)) or r[i - 1] is None for r in rows[:200]):
            for cell in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell[0].number_format = "#,##0.00" if any(isinstance(r[i - 1], float) for r in rows[:200]) else "#,##0"
    ws.freeze_panes = "A2"
    if title or sql or meta:
        about = wb.create_sheet("About")
        about.append(["Title", title])
        for k, v in meta.items():
            about.append([k, str(v)])
        if sql:
            about.append(["SQL", sql])
        about.column_dimensions["A"].width = 14
        about.column_dimensions["B"].width = 100
        for row in about.iter_rows():
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _xlsx_value(v: Any) -> Any:
    if v is None or isinstance(v, (int, float, str)):
        return v
    return str(v)


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:,.2f}"
    return str(v)


def conversation_markdown(turns: list[dict[str, Any]], title: str = "BI Copilot conversation") -> str:
    lines = [f"# {title}", ""]
    for t in turns:
        role = "**You**" if t.get("role") == "user" else "**Copilot**"
        lines += [f"{role}: {t.get('text', '')}", ""]
        r = t.get("result") or {}
        if r.get("sql"):
            lines += ["```sql", r["sql"], "```", ""]
        if r.get("columns") and r.get("rows"):
            lines.append("| " + " | ".join(r["columns"]) + " |")
            lines.append("|" + "---|" * len(r["columns"]))
            for row in r["rows"][:50]:
                lines.append("| " + " | ".join(_cell(v) for v in row) + " |")
            if len(r["rows"]) > 50:
                lines.append(f"... {len(r['rows'])} rows")
            lines.append("")
    return "\n".join(lines) + "\n"
